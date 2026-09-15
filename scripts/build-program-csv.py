import openpyxl, csv, re, io
from openpyxl.utils import column_index_from_string as ci
BASE="/Users/andrew/Dropbox/01_active/every*where/2026/"
wb=openpyxl.load_workbook(BASE+"program/ADW PROGRAM DETAILS -  Updated 14.09.xlsx",data_only=True)
ws=wb["PROGRAM"]; tk=wb["TICKET + IMAGE LINKS"]

def g(r,c,keep_lines=False):
    v=ws.cell(row=r,column=ci(c)).value
    if v is None: return ""
    s=str(v)
    # Newlines carry meaning in the contributors cell, so only spaces are collapsed
    # there. Everywhere else a cell is a single run of text.
    return re.sub(r'[ \t]+',' ',s).strip() if keep_lines else re.sub(r'\s+',' ',s).strip()

def unquote(s):
    # A few blurbs open with a quotation mark that never closes, an artefact of
    # pasting into the form. A lone one reads as a typo on the page.
    if s.count('"')==1: s=s.replace('"','').strip()
    return s

def people(raw):
    """One name per line where the cell uses lines, else comma-separated.
    A line ending in ':' is a heading ("Panelists:"), not a person."""
    parts=raw.split("\n") if "\n" in raw else re.split(r',(?![^(]*\))', raw)
    return [p.strip(" ,") for p in parts if p.strip(" ,") and not p.strip().endswith(":")]

# Category comes from the section blocks: column K labels a block, blank rows separate them.
SECTIONS=[(12,59,"EXH"),(61,64,"INST"),(66,72,"OPEN"),(74,74,"CONV"),(76,94,"CONV"),
          (96,100,"TOUR"),(102,105,"WORK")]
DAYS=[("Q","thu 8 oct"),("R","tue 13 oct"),("S","wed 14 oct"),("T","thu 15 oct"),
      ("U","fri 16 oct"),("V","sat 17 oct"),("W","sun 18 oct")]

# Ticket sheet rows, keyed by their own event name.
TICK={}
for r in range(3,tk.max_row+1):
    n=tk.cell(row=r,column=7).value
    if n: TICK[str(n).strip()]=(tk.cell(row=r,column=14).value, tk.cell(row=r,column=15).value)

# Hand-checked title map: PROGRAM event -> TICKET sheet event. Fuzzy matching produced
# wrong pairs (SLOT->soft, EST OPEN STUDIO->Mixed Goods), so every pair here is explicit.
MAP={
 "10 x 10":"10x10","A ROOM OF ONE'S OWN":"A Room of One’s Own",
 "ADW CLOSING PARTY x COMPANY WORKS":"ADW CLOSING PARTY x COMPANY WORKS",
 "AFTER HOURS":"After Hours","AFTERLIFE":"Afterlife","ANTHO":"ANTHO","AT THE TABLE":"At The Table",
 "BENCHED":"BENCHED","BLOOM":"BLOOM","CARDENING CLUB":"Cardening Club","DESIGN FOR CARE":"Design for Care",
 "GROUNDED":"Grounded","LIVING ON THE CEILING":"Living on the ceiling","MADE IN AUS":"MADE IN AUS",
 "MAKING SPACE":"making space","MOVEMENT ACROSS SCALES":"Movement Across Scales",
 "NEO TEEN TAKEOVER : LETS GOGH (13-17 years)":"Neo Teen Takeover: Let's Gogh",
 "NOW: NEW/OLD WORKS":"NOW: New/Old Works","OBJECT TO MATERIAL":"Object to Material",
 "PUBLIC NOTICE":"Public Notice","POWER SURGE: INTO THE GRID":"Power Surge: Into The Grid",
 "PROCESSED":"PROCESSED","SHOP GENERATOR":"SHOP GENERATOR","SKANGAROOVIA":"Skangaroovia","SOFT":"soft",
 "STARRY STARRY NIGHT":"Starry Starry Night","STEEL SKIN (performance)":"STEEL SKIN",
 "SUPER PARTY":"Super Party","TAC COLLECTION":"taC Collection","TO FEEL IN PLACE":"To Feel In Place",
 "TOKEN":"TOKEN","TOMORROW TALENT":"Tomorrow Talent","TORPOR DAY BED":"Torpor day bed",
 "TWO WAY":"Two-way","UNBUILT":"UNBUILT","UNDER A BRIDGE":"Under a Bridge",
 "WASTE RECONSIDERED":"WASTE RECONSIDERED","X KEULEMANS":"X Guy Keulemans",
 "YOKE, OBJECTS FOR HUMANS":"YOKE • objects for human","MIXED GOODS : OPEN STUDIO":"Mixed Goods Open Studio",
 "YADURA : OPEN HOUSE":"Yadura","FROM POLICY TO PLACE":"From Polict to Place",
 "BESPOKE INTENTION":"Bespoke Intention","COVERS AND PAGES":"Covers + Pages: A publishing design chat",
 "DESIGNING DEEPLY":"Designing Deeply – The paradox of universal design",
 "ÉMIGRÉ : SPEAKER SERIES":"Émigré: Speaker Series at Lynton Residence",
 "In Conversation: Insights, Process and Material Integrity in Lighting Design":
   "In Conversation: Insights, Process and Material Integrity in Lighting Design ",
 "JAM FACTORY x GUY KEULEMANS":"JamFactory Wednesday Talk: Guy Keulemans ",
 "JON GOULDER x AGSA":"Tuesday Talk - Jon Goulder in conversation with Rebecca Evans",
 "MADE HERE":"Made Here","METIS (Navigating Systems Change)":"METIS [ navigating systems change ]",
 "MONET TO MATISSE : ART TO EXHIBITION":"Art of Exhibition: Behind the scenes of Monet to Matisse",
 "THE FRINGE CONDITION : EXTREME DESIGN":"The Fringe Condition: Designing for Life at the Extremes",
 "YES YOU CAN DO A PHD":"Yep, you can do a PhD in design!","YOU'RE NOT FINISHED YET":"You're Not Finished Yet!",
 "EVERY*BIKE TOUR":"Every*bike Tour","LISTENING WITH ARCHITECTURE":"Listening with Architecture",
 "MODERNIST ADELAIDE WALKING TOUR":"Modernist Adelaide walking tour",
 "TORRENS RIVER CITY WALK":"Torrens River City walk","THE DIGITAL PAUSE - Collage Evening":"The Digital Pause",
 "THREADING STORIES":"Threading Stories","WAX CARVING : RING MAKING":"wax carving ring making class ",
 "INSIDE INNOVATION : INSIDE QANTAS TOUR":"Inside Innovation: A Walk Through Qantas’ New Product Innovation Centre",
}
TICKN={k.strip():v for k,v in TICK.items()}
FREE=re.compile(r'^(free|0|none|nil|\$?0(\.00)?)\b',re.I)

def ticket(title):
    name=MAP.get(title)
    if name is None: return "","",None
    price,link=TICK.get(name, TICKN.get(name.strip(),(None,None)))
    link=str(link).strip() if link else ""
    if not re.match(r'^(https?://|www\.)',link): link=""      # "TBC@TBC.com" etc. are not links
    # A bare "www.example.com" in an href is read as a page on this site, so the
    # button loops back to the program instead of leaving it.
    if link.startswith("www."): link = "https://" + link
    if link.startswith("http://"): link = "https://" + link[7:]
    # console.humanitix.com is the organiser's back end - a visitor gets a login
    # screen, not a ticket page. The sheet had one of these.
    if "console.humanitix.com" in link: link = ""
    priced = bool(price) and not FREE.match(str(price).strip())
    return ("yes" if (link or priced) else ""), link, price

rows=[]; notes=[]
for lo,hi,cat in SECTIONS:
    for r in range(lo,hi+1):
        title=g(r,'L')
        if not title: continue
        sessions=[(lab,g(r,col)) for col,lab in DAYS]
        tickd,link,price=ticket(title)
        contributors="; ".join(people(g(r,'AI',keep_lines=True)))
        rows.append({"category":cat,"title":title,"ticketed":tickd,"venue":g(r,'M'),
                     "blurb":unquote(g(r,'AK')),"link":link,"socials":"","note":"","contributors":contributors,"adw_presented":"",
                     **{lab:val for lab,val in sessions},"_row":r,"_price":price})

# SHOPFRONTS is entered twice (row 42 exhibition, row 63 installation) with the same
# all-day times; one row carries a Sunday tour. The site needs one row per title.
# The sheet abbreviates it to SHOPFRONTS; its published name, on Humanitix and in
# the team's own copy, is the Shopfront Design Circuit.
for x in rows:
    if x["title"]=="SHOPFRONTS": x["title"]="SHOPFRONT DESIGN CIRCUIT"

sf=[x for x in rows if x["title"]=="SHOPFRONT DESIGN CIRCUIT"]
if len(sf)==2:
    keep,drop=sf[0],sf[1]
    for _,lab in DAYS:
        vals=[v for v in (keep[lab],drop[lab]) if v]
        # An all-day opening reads first, with the tour as the second session.
        vals.sort(key=lambda v: 0 if v.lower().startswith("(all day") else 1)
        keep[lab]="; ".join(dict.fromkeys(vals))
    rows.remove(drop)
    notes.append(f"SHOPFRONT DESIGN CIRCUIT appeared twice (rows {keep['_row']} and {drop['_row']}); merged into one row, "
                 f"sun 18 oct = '{keep['sun 18 oct']}'.")

for x in rows:
    if not any(x[lab] for _,lab in DAYS):
        notes.append(f"row {x['_row']} '{x['title']}' has no times in any day column.")

# The ADW mark is site metadata the spreadsheet has no column for, so it is carried
# across by title rather than read from the sheet. SHOPFRONTS was "ADW X SHOPFRONT
# DESIGN CIRCUIT" when it was flagged.
ADW={"ADW OPENING PARTY x SHORT NOTICE","ADW CLOSING PARTY x COMPANY WORKS",
     "SHOPFRONT DESIGN CIRCUIT"}
for x in rows:
    if x["title"] in ADW: x["adw_presented"]="yes"

# The Shopfront Design Circuit never went through the EOI form, so neither sheet
# carries anything for it beyond the dates and times. Its details were supplied
# directly by the program team (14 Sep 2026) and are held here so that
# regenerating from a later spreadsheet does not wipe them.
# ADW's 15 Sep pass over the live program. Ticketing is per day now, so an
# exhibition that is free all week with one ticketed opening says so.
OVERRIDES={
 # The sheet gave a console.humanitix.com link - the organiser's own admin page.
 # Dropped until ADW supplies the public one.
 "GROUNDED":{"ticketed":"wed 14 oct", "note":"Registration link released soon"},
 "SUPER PARTY":{"note":"Registration link released soon"},
 "LEARNING FROM YITPI YARTAPUULTIKU":{"note":"Registration link released soon"},
 # The sheet left a note where the description belongs; tickets are not released.
 "THE AUSTRALIAN (DESIGN) DREAM":{"blurb":"", "note":"Registration link released soon"},
 "DESIGNING WITH COUNTRY: Conversations on Place, Practice and Responsibility":{
   "category":"CONV", "note":"Registration link released soon"},
 "ADW OPENING PARTY x SHORT NOTICE":{
   # The sheet left "add short notice description" where the blurb belongs.
   "blurb":"",
   "note":"Registration link released soon. WIN: Ukiyo Residency door prize, announced by the Honourable "
          "Lord Mayor Jane Lomax-Smith"},
 # Supplied by the SLOT team via Hannah, 14 Sep 2026. The Instagram accounts for
 # the host, curator and supporter have nowhere to go: socials is one field per
 # event, and contributors are plain names. Recorded here so they are not lost -
 # SODA Objects @sodaobjects, Table Wines @table.wines, Bronwyn Marshall
 # @marshallstudio_.
 "SLOT":{
  "ticketed":"thu 15 oct",
  "thu 15 oct":"opening 4pm - 7pm",
  "link":"https://events.humanitix.com/slot-opening",
  "socials":"@slot_adw",
  "contributors":"; ".join([
    "Curated by Bronwyn Marshall",
    "Hosted by SODA Objects",
    "Supported by Table Wines",
    "Designers & makers to be revealed"]),
 },
 "CO-DESIGNING YITPI YARTAPUULTIKU":{
  "venue":"Allan Scott Auditorium H2-16, Fenn Place, Adelaide University, ADL CBD",
  "ticketed":"yes",
  "contributors":"Ashley Halliday (Ashley Halliday Architects); Warwick Keates (WAX Design)",
  "blurb":"\n\n".join([
   "YITPI YARTAPUULTIKU - Soul of Port Adelaide: The Story Behind an Award-Winning Place",
   "Discover the remarkable story behind YITPI YARTAPUULTIKU - Soul of Port Adelaide, a "
   "multi-award-winning public realm project that has transformed the heart of Port Adelaide. "
   "Join Architect Ashley Halliday and Landscape Architect Warwick Keates as they take the "
   "audience behind the scenes of the project's conception, design development, and delivery.",
   "This engaging presentation explores how architecture, landscape architecture, planning, "
   "culture, history, and community aspirations were woven together to create a place of "
   "lasting significance. Through insights into the co-design process, working with cultural "
   "knowledge, key challenges, and moments of innovation, Ashley and Warwick reveal the "
   "thinking that shaped a project now recognised for its design excellence and contribution "
   "to public life.",
   "Whether you are a design professional, student, or community member, this presentation "
   "offers a unique opportunity to understand how thoughtful, place-based design can celebrate "
   "identity, strengthen connection, and leave a lasting legacy for future generations.",
   "A free event, bookings essential. Presented by the School of Architecture and Built "
   "Environment, Adelaide University."]),
 },
 "SHOPFRONT DESIGN CIRCUIT":{
 "venue":"Various Locations, East End, ADL CBD",
 "link":"https://events.humanitix.com/shopfront-design-circuit-tour-and-adw-farewell",
 # Only the Sunday walking tour is ticketed; the shopfronts are 24/7 and free.
 "ticketed":"sun 18 oct",
 "note":"WIN: Aesop People's Choice Prize - visit to vote, vote to win",
 "blurb":("The annual Shopfront design challenge sees the East End buzz with a number of "
          "business and designer collaborations for 24/7 viewing. This multi-venue exhibition "
          "has seen local businesses hand over their shopfronts for quick thinking design "
          "responses from selected Adelaide based designers and architects. For Adelaide "
          "Design Week 2026, we have expanded again - with more venues in the East End - "
          "bringing an intimate and broad range of responses. Join us for a dynamic walking "
          "tour of the Shopfront Design Circuit exhibitions, finishing at Honeydripper for an "
          "afternoon to farewell Adelaide Design Week. All are welcome, with additional "
          "collaborations and the final walking tour route to be announced."),
 "socials":"@shopfront_design_sprint",
 "contributors":"; ".join([
   "Jewel of Thought Records x Tom Borgas",
   "Galeria Grafika x Tristan Kerr",
   "Utopian Creations x Jake Lane",
   "Naomi Murrel x Sophia Moore",
   "Shop Pond / Filter Store x Moraene Architecture & Design Collective",
   "Nudie Jeans x Will Cheeseman and Oliver Hyde",
   "Miss Gladys Sim Choon x Claire Markwick-Smith",
   "Aesop x Andrew Carvolth"]),
 },
}
for x in rows:
    if x["title"]=="DESIGNING YITPI": x["title"]="CO-DESIGNING YITPI YARTAPUULTIKU"

for x in rows:
    x.update(OVERRIDES.get(x["title"],{}))

HDR=["category","title","ticketed","venue","blurb","link","socials","note","contributors","adw_presented"]+[l for _,l in DAYS]
out=io.StringIO()
w=csv.DictWriter(out,fieldnames=HDR,extrasaction="ignore",lineterminator="\n")
w.writeheader()
for x in rows: w.writerow(x)
open(BASE+"adw-website/src/_data/program-2026.csv","w").write(out.getvalue())

print(f"{len(rows)} events written")
used={MAP[x['title']] for x in rows if x['title'] in MAP}
print("\nTicket rows not used:", *sorted(set(TICK)-used), sep="\n  ")
print("\nEvents with no ticket-sheet match:", *[x['title'] for x in rows if x['title'] not in MAP], sep="\n  ")
print("\nNotes:", *notes, sep="\n  ")
